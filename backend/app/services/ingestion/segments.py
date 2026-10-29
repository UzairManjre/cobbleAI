"""Structured extraction: course files -> located text segments.

Unlike pdf_extractor.extract_text (one flat string), every Segment keeps where
its text came from (page, slide, or lecture time range), so chunks, citations
and graph provenance can point back to the exact place in the source.
"""

from dataclasses import dataclass, field, asdict
from io import BytesIO
from typing import List, Optional


@dataclass
class Segment:
    text: str
    source_type: str  # pdf_page | slide | docx_para | transcript | frame_ocr | chapter_summary
    page: Optional[int] = None  # 1-based PDF page
    slide: Optional[int] = None  # 1-based slide number
    t_start: Optional[float] = None  # seconds into a lecture
    t_end: Optional[float] = None
    section_title: Optional[str] = None
    extra: dict = field(default_factory=dict)

    def to_dict(self) -> dict:
        return asdict(self)


def extract_pdf_segments(file_bytes: bytes) -> List[Segment]:
    import pdfplumber
    segments = []
    with pdfplumber.open(BytesIO(file_bytes)) as pdf:
        for number, page in enumerate(pdf.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                first_line = text.splitlines()[0].strip()
                segments.append(Segment(
                    text=text,
                    source_type="pdf_page",
                    page=number,
                    section_title=first_line[:120] or None,
                ))
    return segments


def _shape_text(shape) -> str:
    parts = []
    if getattr(shape, "has_text_frame", False) or hasattr(shape, "text"):
        text = getattr(shape, "text", "") or ""
        if text.strip():
            parts.append(text.strip())
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_pptx_segments(file_bytes: bytes) -> List[Segment]:
    import pptx
    presentation = pptx.Presentation(BytesIO(file_bytes))
    segments = []
    for number, slide in enumerate(presentation.slides, start=1):
        title = None
        try:
            if slide.shapes.title is not None and slide.shapes.title.text.strip():
                title = slide.shapes.title.text.strip()
        except Exception:
            title = None
        body = [t for t in (_shape_text(s) for s in slide.shapes) if t]
        notes = ""
        try:
            if slide.has_notes_slide:
                notes = (slide.notes_slide.notes_text_frame.text or "").strip()
        except Exception:
            notes = ""
        text = "\n".join(body)
        if notes:
            text = f"{text}\nSpeaker notes: {notes}" if text else f"Speaker notes: {notes}"
        if text.strip():
            segments.append(Segment(text=text.strip(), source_type="slide", slide=number, section_title=title))
    return segments


def extract_docx_segments(file_bytes: bytes) -> List[Segment]:
    import docx
    document = docx.Document(BytesIO(file_bytes))
    segments = []
    heading = None
    buffer: List[str] = []

    def flush():
        if buffer:
            segments.append(Segment(text="\n".join(buffer), source_type="docx_para", section_title=heading))
            buffer.clear()

    for para in document.paragraphs:
        text = (para.text or "").strip()
        if not text:
            continue
        style = (para.style.name if para.style is not None else "") or ""
        if style.lower().startswith("heading") or style.lower() == "title":
            flush()
            heading = text[:120]
            buffer.append(text)
        else:
            buffer.append(text)
    flush()

    for table in document.tables:
        rows = []
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            segments.append(Segment(text="\n".join(rows), source_type="docx_para", section_title="Table"))
    return segments


def extract_segments(file_type: str, file_bytes: bytes) -> List[Segment]:
    if file_type == "pdf":
        return extract_pdf_segments(file_bytes)
    if file_type == "pptx":
        return extract_pptx_segments(file_bytes)
    if file_type == "docx":
        return extract_docx_segments(file_bytes)
    raise ValueError("Unsupported file type")


def segments_to_text(segments: List[Segment]) -> str:
    return "\n".join(s.text for s in segments)
